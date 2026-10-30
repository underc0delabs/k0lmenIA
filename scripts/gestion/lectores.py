"""Lectura de insumos locales, independiente de la herramienta de gestión.

- Casos manuales desde la planilla .xlsx de generar_casos.py.
- Escenarios BDD desde un .feature.
- Resultados de ejecución desde el cucumber-report.json de k0lmena (web, api o mobile),
  con sus evidencias (capturas embebidas y videos).
"""
import base64
import hashlib
import re
from pathlib import Path

from comun import ErrorGestion, necesitar, normalizar_estado

PRIORIDADES = {"crítica": "Crítica", "critica": "Crítica", "alta": "Alta", "media": "Media", "baja": "Baja"}


def prioridad_de_tags(tags):
    """La prioridad de un escenario BDD desde sus tags: un tag explícito (@critica/@alta/@media/@baja)
    manda; si no hay, @smoke vale Alta; si no, Media."""
    for t in tags:
        p = PRIORIDADES.get(str(t).lower())
        if p:
            return p
    return "Alta" if any(str(t).lower() == "smoke" for t in tags) else "Media"


def _pasos_desde_texto(pasos, datos, resultado):
    """'1. Ir a...\\n2. Completar...' -> [{accion, datos, resultado}]. El resultado esperado del
    caso va en el último paso y los datos de prueba en el primero."""
    lineas = [re.sub(r"^\s*\d+[.)-]\s*", "", l).strip() for l in (pasos or "").splitlines()]
    lineas = [l for l in lineas if l]
    if not lineas:
        lineas = ["(sin pasos)"]
    salida = [{"accion": l, "datos": "", "resultado": ""} for l in lineas]
    salida[0]["datos"] = datos or ""
    salida[-1]["resultado"] = resultado or ""
    return salida


def leer_xlsx(ruta):
    """Lee la planilla de casos (formato de scripts/generar_casos.py)."""
    necesitar("openpyxl")
    from openpyxl import load_workbook

    ws = load_workbook(ruta, data_only=True).active
    encabezados, fila_enc, modulo = None, None, ""
    for fila in ws.iter_rows(min_row=1, max_row=15):
        valores = [str(c.value or "").strip() for c in fila]
        if valores and valores[0].startswith("Módulo:"):
            modulo = valores[0].split(":", 1)[1].strip()  # "HU-001 — Inicio de sesión" (fila 2 de la planilla)
        if "ID #" in valores and "Título" in valores:
            encabezados, fila_enc = valores, fila[0].row
            break
    if not encabezados:
        raise ErrorGestion(f"{ruta}: no encuentro la fila de encabezados (ID #, Título, ...).")
    col = {nombre: i for i, nombre in enumerate(encabezados)}

    def val(fila, nombre):
        i = col.get(nombre)
        return str(fila[i] or "").strip() if i is not None and i < len(fila) else ""

    casos = []
    for fila in ws.iter_rows(min_row=fila_enc + 1, values_only=True):
        cid = val(fila, "ID #")
        if not re.match(r"^CP(-API)?-\d+", cid):
            continue
        casos.append({
            "id": cid,
            "titulo": val(fila, "Título"),
            "descripcion": val(fila, "Descripción"),
            "precondiciones": val(fila, "Precondiciones"),
            "prioridad": PRIORIDADES.get(val(fila, "Prioridad").lower(), "Media"),
            "etiquetas": [e.lstrip("@") for e in re.split(r"[\s,]+", val(fila, "#Etiquetas")) if e],
            "tipo": "manual", "modulo": modulo,
            "tiempo_min": val(fila, "Tiempo estimado (min)"),
            "pasos": _pasos_desde_texto(val(fila, "Pasos"), val(fila, "Datos de Prueba"),
                                         val(fila, "Resultado Esperado") or val(fila, "Resultado esperado")),
        })
    if not casos:
        raise ErrorGestion(f"{ruta}: no encontré casos con ID CP-XXX.")
    return casos


_INICIO_ESCENARIO = re.compile(r"^\s*(Scenario Outline|Scenario Template|Scenario|Example|Rule):")
_INICIO_BLOQUE = re.compile(r"^\s*(Scenario Outline|Scenario Template|Scenario|Example|Rule|Background|Feature):")


def _fin_de_bloque(lineas, j):
    """True si en la línea j empieza otro escenario, regla o feature. Un tag (@x) cierra el bloque
    solo si lo que sigue (después de más tags y comentarios) es un escenario o una regla: los tags
    de un "Examples:" son parte del Scenario Outline."""
    l = lineas[j].strip()
    if _INICIO_BLOQUE.match(l):
        return True
    if l.startswith("@"):
        k = j + 1
        while k < len(lineas) and (not lineas[k].strip() or lineas[k].strip().startswith(("@", "#"))):
            k += 1
        return k < len(lineas) and bool(_INICIO_BLOQUE.match(lineas[k].strip()))
    return False


def leer_feature(ruta):
    """Escenarios de un .feature como casos de tipo 'cucumber' (gherkin completo del escenario).
    El ID del caso sale del tag @CP-XXX del escenario; si no tiene, se usa el título.
    Soporta Background (de la feature y de cada Rule), Scenario Outline con varios Examples
    (también con tags), Example/Scenario Template y Rule."""
    texto = Path(ruta).read_text(encoding="utf-8-sig")
    lineas = texto.splitlines()
    background, background_regla, casos, tags_feature, tags_regla = [], [], [], [], []
    i = 0
    tags_pend = []
    en_regla = False
    feature = ""
    while i < len(lineas):
        l = lineas[i].strip()
        if l.startswith("@"):
            tags_pend += [t.lstrip("@") for t in l.split() if t.startswith("@")]
        elif l.startswith("Feature:"):
            tags_feature, tags_pend = tags_pend, []
            feature = l.split(":", 1)[1].strip()
        elif l.startswith("Rule:"):
            en_regla, tags_regla, tags_pend, background_regla = True, tags_pend, [], []
        elif l.startswith("Background:"):
            destino = background_regla if en_regla else background
            j = i + 1
            while j < len(lineas) and not _fin_de_bloque(lineas, j):
                destino.append(lineas[j].rstrip())
                j += 1
            i = j
            continue
        elif _INICIO_ESCENARIO.match(l):
            titulo = l.split(":", 1)[1].strip()
            cuerpo = [l]
            j = i + 1
            while j < len(lineas) and not _fin_de_bloque(lineas, j):
                cuerpo.append(lineas[j].rstrip())
                j += 1
            while cuerpo and not cuerpo[-1].strip():
                cuerpo.pop()
            tags = tags_pend
            cid = next((t for t in tags if re.match(r"^CP(-API)?-\d+$", t)), titulo)
            fondo = [x for x in background + background_regla if x.strip()]
            gherkin = "\n".join((["Background:"] + fondo + [""] if fondo else []) + cuerpo)
            todas = tags_feature + tags_regla + tags
            casos.append({
                "id": cid, "titulo": titulo, "descripcion": "", "precondiciones": "",
                "prioridad": prioridad_de_tags(todas),
                "etiquetas": list(dict.fromkeys(t for t in todas if not re.match(r"^CP(-API)?-\d+$", t))),
                "tipo": "cucumber", "gherkin": gherkin.strip(), "pasos": [], "feature": feature, "modulo": feature,
            })
            tags_pend = []
            i = j
            continue
        i += 1
    if not casos:
        raise ErrorGestion(f"{ruta}: no encontré escenarios.")
    return casos


def leer_casos(origen, ids=None):
    sufijo = Path(origen).suffix.lower()
    if sufijo == ".xlsx":
        casos = leer_xlsx(origen)
    elif sufijo == ".feature":
        casos = leer_feature(origen)
    else:
        raise ErrorGestion("El origen de los casos tiene que ser un .xlsx (manuales) o un .feature (BDD).")
    if ids:
        pedidos = set(ids)
        casos = [c for c in casos if c["id"] in pedidos]
        faltan = pedidos - {c["id"] for c in casos}
        if faltan:
            raise ErrorGestion(f"No encontré en {origen} los casos: {', '.join(sorted(faltan))}.")
    # Tiempo estimado de ejecución manual (scripts/estimacion.py): viaja a la herramienta de gestión.
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from estimacion import estimar_caso
    for c in casos:
        c["estimacion"] = estimar_caso(c)
    return casos


def leer_resultados(reporte_json, dir_evidencias):
    """cucumber-report.json de k0lmena -> lista de resultados con evidencias en archivos.

    Las capturas embebidas (image/png) se escriben en dir_evidencias. El video se toma del
    adjunto HTML (<video src="videos/...">) que deja el hook. Con reintentos queda el último intento.
    """
    import json

    reporte = Path(reporte_json)
    if not reporte.exists():
        raise ErrorGestion(f"No existe {reporte}. Corré primero npm test en herramientas/k0lmena.")
    base_reporte = reporte.parent
    dir_evidencias = Path(dir_evidencias)
    dir_evidencias.mkdir(parents=True, exist_ok=True)
    features = json.loads(reporte.read_text(encoding="utf-8"))

    por_escenario = {}
    for feature in features:
        for esc in feature.get("elements", []):
            if esc.get("type") == "background":
                continue
            tags = list(dict.fromkeys(t["name"].lstrip("@") for t in esc.get("tags", [])))
            estados, error, capturas, videos, gifs, textos = [], "", [], [], [], []
            duracion_ns = 0
            for paso in esc.get("steps", []):
                res = paso.get("result", {})
                estados.append(res.get("status", "undefined"))
                duracion_ns += int(res.get("duration") or 0)  # Cucumber la da en nanosegundos
                if res.get("status") == "failed" and not error:
                    error = (res.get("error_message") or "").strip()
                for emb in paso.get("embeddings", []):
                    mime = emb.get("mime_type", "")
                    if mime == "image/png":
                        capturas.append(base64.b64decode(emb["data"]))
                    elif mime == "text/html":
                        html = base64.b64decode(emb["data"]).decode("utf-8", "replace")
                        videos += re.findall(r'<video[^>]*src="([^"]+)"', html)
                        gifs += re.findall(r'<img[^>]*src="([^"]+[.]gif)"', html)
                    elif mime == "text/plain":
                        textos.append(base64.b64decode(emb["data"]).decode("utf-8", "replace"))
            bloqueo_antibots = any(t.startswith("BLOQUEO ANTI-BOTS") for t in textos)
            if ("failed" in estados or "ambiguous" in estados) and bloqueo_antibots:
                # El hook de k0lmena (tools/navegador) detectó que la protección anti-bots frenó al navegador:
                # no es un resultado de la app, así que va como bloqueado y con el motivo como error.
                estado = "bloqueado"
                error = next(t for t in textos if t.startswith("BLOQUEO ANTI-BOTS"))
            elif "failed" in estados or "ambiguous" in estados:
                estado = "fallido"
            elif all(e == "passed" for e in estados):
                estado = "aprobado"
            else:
                estado = "pendiente"
            clave = f'{feature.get("uri", "")}:{esc.get("line")}:{esc.get("name")}'
            por_escenario[clave] = {
                "escenario": esc.get("name", ""), "feature": feature.get("name", ""),
                "tags": tags, "estado": normalizar_estado(estado), "error": error,
                "duracion_s": round(duracion_ns / 1e9, 2) if duracion_ns else None,
                "_capturas": capturas, "_videos": videos, "_gifs": gifs, "textos": textos,
            }

    resultados = []
    for r in por_escenario.values():
        nombre = re.sub(r"[^\w.-]+", "_", r["escenario"])[:80]
        evidencias = []
        for n, png in enumerate(r.pop("_capturas")):
            archivo = dir_evidencias / f"{nombre}-{hashlib.md5(png).hexdigest()[:8]}-{n}.png"
            archivo.write_bytes(png)
            evidencias.append({"tipo": "captura", "archivo": str(archivo)})
        for tipo, clave in (("video", "_videos"), ("gif", "_gifs")):
            for src in r.pop(clave):
                # web: "videos/x.webm" (relativo a reports/web) · mobile: "../videos/x.mp4" (relativo a html/features)
                limpio = src.replace("../", "")
                for candidato in (base_reporte / limpio, base_reporte / "html" / limpio):
                    if candidato.exists():
                        evidencias.append({"tipo": tipo, "archivo": str(candidato.resolve())})
                        break
        r["evidencias"] = evidencias
        resultados.append(r)
    return resultados


_TIPOS_EVIDENCIA = {".png": "captura", ".jpg": "captura", ".jpeg": "captura", ".webp": "captura", ".gif": "gif",
                    ".webm": "video", ".mp4": "video", ".mov": "video"}


def _ids_junit(nombres, propiedades):
    """Los IDs que identifican un test JUnit. En el nombre y la clase: CP-001 (también cp_001, CP001 o CP-API-003),
    HU-001 y keys escritas como tag o entre corchetes (@PROJ-45, [TC-12]); así test_cp_001_login de pytest o
    "CP-001 - Login válido" de Cypress quedan identificados, y "UTF-8" en un nombre no pasa por una key.
    En las properties (tags, test_key, test_id…), cualquier key."""
    tags = []
    for texto, es_propiedad in [(n, False) for n in nombres] + [(v, True) for v in propiedades]:
        texto = str(texto or "")
        tags += re.findall(r"(?<![\w.])@([A-Za-z][\w.-]*)", texto)
        tags += re.findall(r"\[([A-Z][A-Z0-9]+-\d+)\]", texto)
        for m in re.finditer(r"(?<![A-Za-z0-9])(cp|hu)[-_ ]?(api[-_ ]?)?(\d+)(?![0-9])", texto, re.I):
            tags.append(m.group(1).upper() + ("-API" if m.group(2) else "") + f"-{int(m.group(3)):03d}")
        if es_propiedad:
            tags += re.findall(r"(?<![A-Za-z0-9-])([A-Z][A-Z0-9]+-\d+)(?![A-Za-z0-9-])", texto)
    return list(dict.fromkeys(tags))


def leer_resultados_junit(origen, dir_evidencias):
    """Reporte JUnit XML (Cypress, Playwright, pytest, Maven/Surefire, Jest, Robot…) -> la misma lista de resultados
    que leer_resultados, para publicar en la herramienta de gestión lo que corrió con otro framework.

    origen: un .xml o una carpeta con varios (Surefire deja uno por clase).
    - El caso sale del nombre del test, de su clase o de sus properties (tags, test_key, test_id…): CP-001,
      cp_001, @CP-001, una key (PROJ-45, TC-12) y la historia (HU-001).
    - failure/error -> fallido, skipped -> pendiente, sin nada -> aprobado. Si un test aparece varias veces
      (reintentos), vale el último.
    - Evidencias: [[ATTACHMENT|ruta]] en system-out/system-err (Playwright, Jenkins), properties attachment y
      testrun_evidence de Xray (base64). Las rutas relativas se buscan desde la carpeta del XML.
    """
    import xml.etree.ElementTree as ET

    origen = Path(origen)
    archivos = sorted(origen.glob("**/*.xml")) if origen.is_dir() else [origen]
    if not archivos or not all(a.exists() for a in archivos):
        raise ErrorGestion(f"No encontré reportes JUnit en {origen}.")
    dir_evidencias = Path(dir_evidencias)
    dir_evidencias.mkdir(parents=True, exist_ok=True)

    por_test = {}
    for archivo in archivos:
        try:
            raiz = ET.parse(archivo).getroot()
        except ET.ParseError as e:
            raise ErrorGestion(f"{archivo}: no es un XML válido ({e}).")
        suites = [raiz] if raiz.tag == "testsuite" else list(raiz.iter("testsuite"))
        if raiz.tag not in ("testsuite", "testsuites"):
            raise ErrorGestion(f"{archivo}: no parece un reporte JUnit (la raíz es <{raiz.tag}>).")
        for suite in suites:
            for tc in suite.findall("testcase"):
                nombre, clase = tc.get("name", ""), tc.get("classname", "")
                props, evidencias = {}, []
                for p in tc.iter("property"):
                    if p.get("name") == "testrun_evidence":
                        for item in p.findall("item"):
                            evidencias.append(("xray", item.get("name") or "evidencia.png", item.text or ""))
                    elif p.get("name") == "attachment":
                        evidencias.append(("ruta", p.get("value") or (p.text or "").strip(), None))
                    else:
                        props[p.get("name", "")] = p.get("value") or (p.text or "")
                salida = "\n".join((x.text or "") for x in tc if x.tag in ("system-out", "system-err"))
                evidencias += [("ruta", r.strip(), None) for r in re.findall(r"\[\[ATTACHMENT\|([^\]]+)\]\]", salida)]
                falla = next((x for x in tc if x.tag in ("failure", "error")), None)
                if falla is not None:
                    estado = "fallido"
                    error = "\n".join(t for t in (falla.get("message", ""), (falla.text or "").strip()) if t)
                elif tc.find("skipped") is not None:
                    estado, error = "pendiente", tc.find("skipped").get("message", "")
                else:
                    estado, error = "aprobado", ""
                try:
                    duracion = round(float(tc.get("time") or 0), 2) or None
                except ValueError:
                    duracion = None
                por_test[(clase, nombre)] = {
                    "escenario": nombre, "feature": clase or suite.get("name", ""),
                    "tags": _ids_junit([nombre, clase], props.values()),
                    "estado": normalizar_estado(estado), "error": error.strip(), "duracion_s": duracion,
                    "textos": [], "origen": "JUnit", "_evidencias": evidencias, "_base": archivo.parent,
                }

    resultados = []
    for r in por_test.values():
        nombre = re.sub(r"[^\w.-]+", "_", r["escenario"])[:80]
        base = r.pop("_base")
        evidencias = []
        for clase, valor, datos in r.pop("_evidencias"):
            if clase == "xray":
                contenido = base64.b64decode(datos)
                archivo = dir_evidencias / f"{nombre}-{hashlib.md5(contenido).hexdigest()[:8]}-{Path(valor).name}"
                archivo.write_bytes(contenido)
            else:
                archivo = Path(valor) if Path(valor).is_absolute() else base / valor
                if not archivo.exists():
                    continue
            evidencias.append({"tipo": _TIPOS_EVIDENCIA.get(archivo.suffix.lower(), "archivo"),
                               "archivo": str(Path(archivo).resolve())})
        r["evidencias"] = evidencias
        resultados.append(r)
    return resultados
