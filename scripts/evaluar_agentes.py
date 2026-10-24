#!/usr/bin/env python3
"""
Evals de los agentes de k0lmenIA: miden la calidad de lo que GENERAN los agentes (no solo los scripts).

Cada caso de evals/casos/<id>/caso.json dice qué agente se evalúa, el pedido en lenguaje natural, los
insumos que se copian a un input/ aislado y los chequeos determinísticos sobre lo que el agente deja
(archivos que tienen que existir, textos que tienen que aparecer o no, cantidad de casos por tipo,
columnas de la planilla, Gherkin válido, IDs de Falta información…). Ver evals/README.md.

Uso (desde la raíz del repo):
    python scripts/evaluar_agentes.py --listar
    python scripts/evaluar_agentes.py --preparar [--caso ID | --agente NOMBRE] [--trabajo DIR]
    python scripts/evaluar_agentes.py --correr   [--caso ID | --agente NOMBRE] [--modelo M] [--presupuesto-usd 2] [--juez]
    python scripts/evaluar_agentes.py --verificar --trabajo DIR
    python scripts/evaluar_agentes.py --solo-verificar CARPETA --caso ID

--preparar   arma un workspace aislado por caso (una copia mínima del repo: CLAUDE.md, REFERENCIA.md, .claude/, plantillas/,
             scripts/, con input/ y output/ vacíos + los insumos del caso), fuera del repo (por defecto en la
             carpeta temporal del sistema), así el agente nunca toca tu input/ ni tu output/.
--correr     prepara, corre el agente con Claude Code en modo headless (claude -p, permisos acotados al
             workspace, sin conectores MCP, con un tope de gasto por caso: --max-budget-usd) y verifica. GASTA TOKENS: cada caso es una conversación entera.
--verificar  verifica los workspaces de una corrida ya preparada o corrida (--trabajo).
--solo-verificar  verifica salidas que ya existen en CARPETA (por ejemplo, la raíz del repo después de
             pedirle el caso a mano al agente). Para lo que el agente responde en el chat, guardá la respuesta
             en CARPETA/_eval/respuesta.txt.
--juez       (opcional) además de los chequeos, le pide a Claude que puntúe la rúbrica del caso. Es solo
             informativo: no cambia el puntaje.

Deja evals/resultados/<fecha-hora>/resultados.json e informe-evals.html (puntaje por caso y por agente,
qué chequeo falló, costo y turnos de cada corrida, y la comparación con la corrida anterior para detectar
regresiones al cambiar un agente o un skill). evals/resultados/ no se versiona.
"""
import argparse
import datetime as dt
import glob
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
EVALS = RAIZ / "evals"
CASOS = EVALS / "casos"
RESULTADOS = EVALS / "resultados"
EVAL_DIR = "_eval"  # dentro de cada workspace: la respuesta y los datos de la corrida de claude

# Lo que se copia del repo al workspace (el resto, como herramientas/ o docs/, no hace falta para evaluar).
COPIAR = ["CLAUDE.md", "AGENTES.md", "REFERENCIA.md", ".claude", "plantillas", "scripts", "requirements.txt"]
IGNORAR_COPIA = shutil.ignore_patterns("__pycache__", "*.pyc", "settings.local.json", "fuentes")
# Herramientas que puede usar el agente en el workspace (claude -p niega todo lo demás).
HERRAMIENTAS = ["Read", "Write", "Edit", "Glob", "Grep", "Agent", "Task", "TodoWrite",
                "Bash(python:*)", "Bash(python3:*)", "Bash(git -C:*)", "Bash(ls:*)"]
PRIORIDADES = ["Crítica", "Alta", "Media", "Baja"]
ESTADOS_CASO = ["N/A", "Pendiente", "En ejecución", "Aprobado", "Fallido", "Bloqueado"]


# ------------------------------------------------------------------ casos
def cargar_casos(caso=None, agente=None, carpeta=CASOS):
    casos = []
    for archivo in sorted(Path(carpeta).glob("*/caso.json")):
        datos = json.loads(archivo.read_text(encoding="utf-8"))
        datos.setdefault("id", archivo.parent.name)
        datos["_dir"] = str(archivo.parent)
        if (caso and datos["id"] != caso) or (agente and datos.get("agente") != agente):
            continue
        casos.append(datos)
    if caso and not casos:
        raise SystemExit(f"No existe el caso {caso!r} en {carpeta}")
    return casos


def faltantes(caso, raiz=RAIZ):
    """Lo que el caso requiere y no está en el repo (por ejemplo, un agente que todavía no existe)."""
    return [r for r in caso.get("requiere", []) if not (Path(raiz) / r).exists()]


# ------------------------------------------------------------------ workspace
def _origen(desde, caso):
    if desde.startswith("repo:"):
        return RAIZ / desde[5:]
    p = EVALS / desde
    return p if p.exists() else Path(caso["_dir"]) / desde


def preparar(caso, destino):
    """Arma el workspace aislado del caso en `destino` y devuelve su ruta."""
    ws = Path(destino)
    if ws.exists():
        shutil.rmtree(ws)
    ws.mkdir(parents=True)
    for nombre in COPIAR:
        src = RAIZ / nombre
        if src.is_dir():
            shutil.copytree(src, ws / nombre, ignore=IGNORAR_COPIA)
        elif src.exists():
            shutil.copy2(src, ws / nombre)
    # input/ y output/ con la misma estructura de carpetas (dos niveles), vacías.
    for base in ("input", "output"):
        for d in [p for p in (RAIZ / base).rglob("*") if p.is_dir() and len(p.relative_to(RAIZ).parts) <= 3] + [RAIZ / base]:
            (ws / d.relative_to(RAIZ)).mkdir(parents=True, exist_ok=True)
    for ins in caso.get("insumos", []):
        src, dst = _origen(ins["desde"], caso), ws / ins["hacia"]
        dst.parent.mkdir(parents=True, exist_ok=True)
        if src.is_dir():
            shutil.copytree(src, dst, dirs_exist_ok=True)
        else:
            shutil.copy2(src, dst)
    env = {k: str(v).replace("{workspace}", ws.as_posix()) for k, v in caso.get("env", {}).items()}
    (ws / ".env").write_text("".join(f"{k}={v}\n" for k, v in env.items()), encoding="utf-8")
    (ws / EVAL_DIR).mkdir()
    (ws / EVAL_DIR / "pedido.txt").write_text(caso["pedido"], encoding="utf-8")
    return ws


def comando_claude(caso, modelo=None, max_turnos=None, ejecutable="claude", presupuesto_usd=None):
    """Línea de claude -p para el caso: headless, salida JSON, permisos acotados, tope de gasto y sin conectores MCP."""
    tope = presupuesto_usd or caso.get("presupuesto_usd", 2)
    cmd = [ejecutable, "-p", caso["pedido"], "--output-format", "json",
           "--permission-mode", "acceptEdits", "--allowedTools", ",".join(HERRAMIENTAS),
           "--max-budget-usd", f"{float(tope):g}", "--strict-mcp-config", "--setting-sources", "project"]
    if max_turnos:  # solo si se pide: no todas las versiones de Claude Code lo aceptan
        cmd += ["--max-turns", str(max_turnos)]
    if modelo:
        cmd += ["--model", modelo]
    return cmd


def correr_claude(caso, ws, modelo=None, max_turnos=None, timeout_s=1800, presupuesto_usd=None):
    """Corre el agente en el workspace y guarda la respuesta y las métricas en _eval/."""
    ejecutable = shutil.which("claude")
    if not ejecutable:
        return {"error": "No encuentro el comando claude (instalá Claude Code)."}
    inicio = dt.datetime.now()
    try:
        r = subprocess.run(comando_claude(caso, modelo, max_turnos, ejecutable, presupuesto_usd), cwd=ws, capture_output=True,
                           text=True, encoding="utf-8", timeout=timeout_s)
    except subprocess.TimeoutExpired:
        return {"error": f"Se cortó a los {timeout_s} s."}
    (Path(ws) / EVAL_DIR / "claude.json").write_text(r.stdout or "", encoding="utf-8")
    try:
        datos = json.loads(r.stdout)
    except ValueError:
        return {"error": (r.stderr or r.stdout or "claude no devolvió JSON").strip()[-500:]}
    (Path(ws) / EVAL_DIR / "respuesta.txt").write_text(str(datos.get("result") or ""), encoding="utf-8")
    return {"costo_usd": datos.get("total_cost_usd"), "turnos": datos.get("num_turns"),
            "duracion_s": round((datos.get("duration_ms") or (dt.datetime.now() - inicio).total_seconds() * 1000) / 1000),
            "error": datos.get("result") if datos.get("is_error") else None}


def juez(caso, ws, modelo=None):
    """Rúbrica opcional evaluada por Claude (informativa). Devuelve {"puntajes": [...]} o {"error": ...}."""
    if not caso.get("rubrica") or not shutil.which("claude"):
        return None
    criterios = "\n".join(f"- {c}" for c in caso["rubrica"])
    pedido = ("Sos un revisor de QA. Evaluá lo que dejó el agente en output/, input/documentacion/ y _eval/respuesta.txt "
              f"para este pedido: {caso['pedido']!r}. Puntuá cada criterio de 0 a 2 (0 = no cumple, 2 = cumple bien):\n"
              f"{criterios}\nRespondé SOLO con JSON: {{\"puntajes\": [{{\"criterio\": \"...\", \"puntaje\": 0, \"motivo\": \"...\"}}]}}")
    cmd = [shutil.which("claude"), "-p", pedido, "--output-format", "json", "--allowedTools", "Read,Glob,Grep",
           "--max-budget-usd", "1", "--strict-mcp-config"] + (["--model", modelo] if modelo else [])
    try:
        r = subprocess.run(cmd, cwd=ws, capture_output=True, text=True, encoding="utf-8", timeout=600)
        texto = json.loads(r.stdout).get("result", "")
        return json.loads(re.search(r"\{.*\}", texto, re.S).group(0))
    except (ValueError, AttributeError, subprocess.TimeoutExpired) as ex:
        return {"error": str(ex)}


# ------------------------------------------------------------------ lectura de salidas
def archivos(raiz, patrones):
    """Archivos de `raiz` que coinciden con los globs (relativos a la raíz)."""
    patrones = [patrones] if isinstance(patrones, str) else patrones
    vistos = []
    for p in patrones:
        for f in sorted(glob.glob(str(Path(raiz) / p), recursive=True)):
            if Path(f).is_file() and f not in vistos:
                vistos.append(f)
    return [Path(f) for f in vistos]


def texto_de(rutas):
    partes = []
    for r in rutas:
        if r.suffix.lower() == ".xlsx":
            partes.append("\n".join(" | ".join(str(v) for v in fila.values()) for fila in leer_xlsx(r)[1]))
        else:
            partes.append(r.read_text(encoding="utf-8", errors="replace"))
    return "\n".join(partes)


def leer_xlsx(ruta):
    """(encabezados, filas, hojas) de la planilla de casos: la fila de encabezados es la que tiene "ID #"."""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True)
    ws = wb.worksheets[0]
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    idx = next((i for i, f in enumerate(filas[:15]) if any(str(c or "").strip() == "ID #" for c in f)), None)
    if idx is None:
        return [], [], {h.title: h for h in wb.worksheets}
    enc = [str(c or "").strip() for c in filas[idx]]
    casos = []
    for f in filas[idx + 1:]:
        if not str(f[0] or "").strip():
            break
        casos.append({enc[i]: ("" if v is None else str(v)) for i, v in enumerate(f) if i < len(enc) and enc[i]})
    return enc, casos, {h.title: h for h in wb.worksheets}


def parsear_gherkin(texto):
    """Estructura mínima de un .feature: {feature, tags_feature, escenarios: [...], errores: [...]}."""
    res = {"feature": None, "tags_feature": [], "escenarios": [], "errores": []}
    tags, actual, en_ejemplos = [], None, False
    for n, linea in enumerate(texto.splitlines(), 1):
        s = linea.strip()
        if not s or s.startswith("#"):
            continue
        if s.startswith("@"):
            tags += s.split()
            continue
        m = re.match(r"(Feature|Background|Scenario Outline|Scenario Template|Scenario|Example|Examples|Rule):\s*(.*)", s)
        if m:
            clave = m.group(1)
            if clave == "Feature":
                if res["feature"] is not None:
                    res["errores"].append(f"Línea {n}: más de un Feature")
                res["feature"], res["tags_feature"] = m.group(2), tags
            elif clave in ("Scenario", "Scenario Outline", "Scenario Template", "Example"):
                actual = {"titulo": m.group(2), "tags": tags, "pasos": [], "outline": clave != "Scenario" and clave != "Example",
                          "ejemplos": 0, "linea": n}
                res["escenarios"].append(actual)
            elif clave == "Examples":
                en_ejemplos = True
                if actual is None or not actual["outline"]:
                    res["errores"].append(f"Línea {n}: Examples fuera de un Scenario Outline")
            elif clave == "Background":
                actual = None
            tags = []
            if clave != "Examples":
                en_ejemplos = False
            continue
        if s.startswith("|"):
            if en_ejemplos and actual:
                actual["ejemplos"] += 1
            continue
        m = re.match(r"(Given|When|Then|And|But|\*)\s+\S", s)
        if m:
            if actual is not None:
                actual["pasos"].append(m.group(1))
            continue
        if s.startswith('"""') or s.startswith("```"):
            continue
        if res["feature"] is None or actual is not None:
            res["errores"].append(f"Línea {n}: no es Gherkin válido ({s[:40]})")
    if res["feature"] is None:
        res["errores"].append("No tiene Feature")
    for esc in res["escenarios"]:
        if "Then" not in esc["pasos"]:
            esc_err = f"El escenario de la línea {esc['linea']} no tiene Then"
            res["errores"].append(esc_err)
        if not ({"Given", "When"} & set(esc["pasos"])):
            res["errores"].append(f"El escenario de la línea {esc['linea']} no tiene Given ni When")
        if esc["outline"] and esc["ejemplos"] < 2:
            res["errores"].append(f"El Scenario Outline de la línea {esc['linea']} no tiene tabla de Examples")
    return res


def _tablas_crudas(texto):
    """Filas de tablas Markdown fuera de bloques de código (formatear_tablas.py no se corrió)."""
    fuera, crudas = True, 0
    for linea in texto.splitlines():
        if linea.strip().startswith("```"):
            fuera = not fuera
            continue
        if fuera and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", linea):
            crudas += 1
    return crudas


# ------------------------------------------------------------------ chequeos
def _r(p):
    return re.compile(p, re.I | re.M | re.S)


def chequear(ch, raiz):
    """Corre un chequeo y devuelve (ok, detalle)."""
    tipo = ch["tipo"]
    rutas = archivos(raiz, ch.get("archivos", [])) if ch.get("archivos") else []
    if tipo == "existe":
        ok = len(rutas) >= ch.get("min", 1)
        return ok, f"{len(rutas)} archivo(s): " + ", ".join(Path(r).name for r in rutas[:5]) if rutas else "No se encontró ningún archivo"
    if ch.get("archivos") and not rutas:
        return False, "No se encontró el archivo a revisar"
    if tipo in ("contiene", "no_contiene"):
        texto = texto_de(rutas)
        hallados = [p for p in ch["patrones"] if _r(p).search(texto)]
        if tipo == "no_contiene":
            return not hallados, ("Aparece: " + "; ".join(hallados)) if hallados else "No aparece"
        if ch.get("modo") == "alguno":
            return bool(hallados), "Coincide: " + "; ".join(hallados) if hallados else "Ningún patrón coincide"
        faltan = [p for p in ch["patrones"] if p not in hallados]
        return not faltan, ("Falta: " + "; ".join(faltan)) if faltan else "Aparecen todos"
    if tipo == "conteo":
        encontrados = _r(ch["patron"]).findall(texto_de(rutas))
        encontrados = [m if isinstance(m, str) else m[0] for m in encontrados]
        n = len(set(x.upper() for x in encontrados)) if ch.get("distintos") else len(encontrados)
        return n >= ch.get("min", 1), f"{n} (mínimo {ch.get('min', 1)})"
    if tipo == "orden":
        texto, pos = texto_de(rutas), 0
        for p in ch["patrones"]:
            m = _r(p).search(texto, pos)
            if not m:
                return False, f"No aparece (o no en orden): {p}"
            pos = m.end()
        return True, "En orden"
    if tipo == "tablas_normalizadas":
        crudas = sum(_tablas_crudas(r.read_text(encoding="utf-8", errors="replace")) for r in rutas)
        return crudas == 0, "Alineadas" if not crudas else f"{crudas} tabla(s) en Markdown crudo (falta formatear_tablas.py)"
    if tipo.startswith("xlsx_"):
        return _chequear_xlsx(ch, rutas[0])
    if tipo == "gherkin":
        return _chequear_gherkin(ch, rutas)
    raise ValueError(f"Tipo de chequeo desconocido: {tipo}")


def _texto_caso(c, claves=("Título", "Descripción", "Datos de Prueba", "#Etiquetas")):
    return " ".join(c.get(k, "") for k in claves)


def _chequear_xlsx(ch, ruta):
    enc, casos, hojas = leer_xlsx(ruta)
    tipo = ch["tipo"]
    if tipo == "xlsx_columnas":
        faltan = [c for c in ch["columnas"] if c not in enc]
        return not faltan, ("Faltan columnas: " + ", ".join(faltan)) if faltan else "Columnas completas"
    if tipo == "xlsx_casos":
        problemas = []
        if len(casos) < ch.get("min", 1):
            problemas.append(f"{len(casos)} casos (mínimo {ch.get('min', 1)})")
        ids = [c.get("ID #", "") for c in casos]
        malos = [i for i in ids if not re.fullmatch(ch.get("ids", r"CP-\d{3}"), i)]
        if malos:
            problemas.append("IDs con formato inválido: " + ", ".join(malos[:5]))
        if len(set(ids)) != len(ids):
            problemas.append("IDs repetidos")
        esperados = [f"CP-{i:03d}" for i in range(1, len(ids) + 1)]
        if ch.get("consecutivos", True) and not malos and ids != esperados:
            problemas.append("IDs no consecutivos desde CP-001")
        prios = [c.get("Prioridad", "") for c in casos]
        if any(p not in PRIORIDADES for p in prios):
            problemas.append("Prioridades fuera de la escala: " + ", ".join(sorted({p for p in prios if p not in PRIORIDADES})))
        elif ch.get("orden_prioridad", True) and [PRIORIDADES.index(p) for p in prios] != sorted(PRIORIDADES.index(p) for p in prios):
            problemas.append("No están ordenados por prioridad (Crítica → Baja)")
        estados = {c.get("Estado", "") for c in casos} - set(ESTADOS_CASO)
        if estados:
            problemas.append("Estados fuera de la escala: " + ", ".join(sorted(estados)))
        sin_resultado = [c.get("ID #") for c in casos if not c.get("Resultado Esperado", "").strip() or not c.get("Pasos", "").strip()]
        if sin_resultado:
            problemas.append("Sin pasos o resultado esperado: " + ", ".join(sin_resultado[:5]))
        return not problemas, "; ".join(problemas) or f"{len(casos)} casos válidos"
    if tipo == "xlsx_tipos":
        detalle, ok = [], True
        for nombre, regla in ch["tipos"].items():
            n = sum(1 for c in casos if _r(regla["patron"]).search(_texto_caso(c)))
            ok &= n >= regla["min"]
            detalle.append(f"{nombre}: {n} (mín. {regla['min']})")
        return ok, " · ".join(detalle)
    if tipo == "xlsx_falta_info":
        marcados = [c for c in casos if "@falta-info" in c.get("#Etiquetas", "") or re.search(r"FI-\d+", c.get("Comentarios", ""))]
        hoja = hojas.get("Falta información")
        items = set()
        if hoja is not None:
            for fila in hoja.iter_rows(values_only=True):
                items |= {m for v in fila for m in re.findall(r"FI-\d+", str(v or ""))}
        ok = len(marcados) >= ch.get("min_casos", 1) and len(items) >= ch.get("min_items", 1)
        return ok, f"{len(marcados)} caso(s) marcados · {len(items)} ítem(s) en la hoja Falta información"
    if tipo == "xlsx_filtro":
        afectados = [c for c in casos if _r(ch["filtro"]).search(_texto_caso(c, ("Título", "Descripción", "Pasos", "Resultado Esperado")))]
        if not afectados:
            return ch.get("ok_si_no_hay", False), "Ningún caso trata el tema"
        malos = []
        for c in afectados:
            marcas = c.get("#Etiquetas", "") + " " + c.get("Comentarios", "")
            if ch.get("debe") and not _r(ch["debe"]).search(marcas):
                malos.append(f"{c.get('ID #')} sin {ch['debe']}")
            if ch.get("no_debe_resultado") and _r(ch["no_debe_resultado"]).search(c.get("Resultado Esperado", "")):
                malos.append(f"{c.get('ID #')} inventa: {c.get('Resultado Esperado', '')[:60]}")
        return not malos, "; ".join(malos) or f"{len(afectados)} caso(s) correctos"
    raise ValueError(f"Tipo de chequeo desconocido: {tipo}")


def _chequear_gherkin(ch, rutas):
    problemas, escenarios, ids, con_fi = [], 0, [], 0
    for r in rutas:
        texto = r.read_text(encoding="utf-8", errors="replace")
        g = parsear_gherkin(texto)
        problemas += [f"{r.name}: {x}" for x in g["errores"]]
        if ch.get("tag_feature") and not any(re.fullmatch(ch["tag_feature"], t) for t in g["tags_feature"]):
            problemas.append(f"{r.name}: el Feature no tiene el tag {ch['tag_feature']}")
        for esc in g["escenarios"]:
            escenarios += 1
            tag_cp = [t for t in esc["tags"] if re.fullmatch(ch.get("tag_escenario", r"@CP-\d{3}"), t)]
            if not tag_cp:
                problemas.append(f"{r.name}: '{esc['titulo'][:40]}' sin tag de caso")
            ids += tag_cp
            if "@falta-info" in esc["tags"] and any(re.fullmatch(r"@FI-\d+", t) for t in esc["tags"]):
                con_fi += 1
        if ch.get("falta_info_min") and not re.search(r"#\s*FALTA INFORMACI[ÓO]N \(FI-\d+\)", texto, re.I):
            problemas.append(f"{r.name}: falta el comentario # FALTA INFORMACIÓN (FI-XX)")
    if len(set(ids)) != len(ids):
        problemas.append("Tags de caso repetidos")
    if escenarios < ch.get("min_escenarios", 1):
        problemas.append(f"{escenarios} escenarios (mínimo {ch.get('min_escenarios', 1)})")
    if con_fi < ch.get("falta_info_min", 0):
        problemas.append(f"{con_fi} escenario(s) con @falta-info @FI-XX (mínimo {ch['falta_info_min']})")
    return not problemas, "; ".join(problemas[:6]) or f"{escenarios} escenarios válidos"


def verificar_caso(caso, raiz):
    """Corre todos los chequeos del caso sobre `raiz` y devuelve el resultado con su puntaje."""
    resultados, total, logrado = [], 0, 0
    for ch in caso.get("chequeos", []):
        try:
            ok, detalle = chequear(ch, raiz)
        except Exception as ex:  # un chequeo roto no corta la evaluación
            ok, detalle = False, f"Error al verificar: {ex}"
        peso = ch.get("peso", 1)
        total += peso
        logrado += peso if ok else 0
        resultados.append({"id": ch["id"], "descripcion": ch.get("descripcion", ch["id"]), "ok": ok, "detalle": detalle, "peso": peso})
    puntaje = round(100 * logrado / total) if total else 0
    estado = "Aprobado" if puntaje >= caso.get("umbral", 80) else "Fallido"
    return {"id": caso["id"], "agente": caso.get("agente", ""), "descripcion": caso.get("descripcion", ""),
            "puntaje": puntaje, "estado": estado, "chequeos": resultados}


# ------------------------------------------------------------------ comparación e informe
def corrida_anterior(destino):
    anteriores = sorted(p for p in Path(destino).parent.glob("*/resultados.json") if p.parent != Path(destino))
    anteriores = [p for p in anteriores if p.parent.name < Path(destino).name]
    return json.loads(anteriores[-1].read_text(encoding="utf-8")) if anteriores else None


def comparar(actual, anterior):
    """Por caso: puntaje antes/después, chequeos que dejaron de pasar (regresiones) y los que ahora pasan."""
    if not anterior:
        return None
    previos = {c["id"]: c for c in anterior.get("casos", [])}
    out = {"fecha_anterior": anterior.get("fecha"), "casos": {}, "regresiones": 0, "mejoras": 0}
    for c in actual["casos"]:
        p = previos.get(c["id"])
        if not p or c["estado"] in ("N/A", "Bloqueado") or p.get("estado") in ("N/A", "Bloqueado"):
            continue
        antes = {x["id"]: x["ok"] for x in p.get("chequeos", [])}
        reg = [x["id"] for x in c["chequeos"] if antes.get(x["id"]) is True and not x["ok"]]
        mej = [x["id"] for x in c["chequeos"] if antes.get(x["id"]) is False and x["ok"]]
        out["casos"][c["id"]] = {"antes": p.get("puntaje"), "despues": c["puntaje"], "regresiones": reg, "mejoras": mej}
        out["regresiones"] += len(reg)
        out["mejoras"] += len(mej)
    return out


def resumen_por_agente(casos):
    agentes = {}
    for c in casos:
        if c["estado"] in ("N/A", "Bloqueado"):
            continue
        a = agentes.setdefault(c["agente"], {"casos": 0, "aprobados": 0, "puntajes": []})
        a["casos"] += 1
        a["aprobados"] += c["estado"] == "Aprobado"
        a["puntajes"].append(c["puntaje"])
    return {k: {"casos": v["casos"], "aprobados": v["aprobados"], "puntaje": round(sum(v["puntajes"]) / len(v["puntajes"]))}
            for k, v in agentes.items()}


def informe_html(datos):
    sys.path.insert(0, str(RAIZ / "scripts"))
    from _estilos_reporte import (e, page, meta, section, kpi_tiles, banner, table, bar_chart, estado_pill, doc_list,
                                  OK, AVISO, CRITICO)
    casos = datos["casos"]
    evaluados = [c for c in casos if c["estado"] in ("Aprobado", "Fallido")]
    aprobados = sum(c["estado"] == "Aprobado" for c in evaluados)
    prom = round(sum(c["puntaje"] for c in evaluados) / len(evaluados)) if evaluados else 0
    comp = datos.get("comparacion") or {}
    costo = sum((c.get("metricas") or {}).get("costo_usd") or 0 for c in casos)
    tono = "ok" if evaluados and aprobados == len(evaluados) and not comp.get("regresiones") else "warn" if prom >= 60 else "bad"
    veredicto = (f"{aprobados} de {len(evaluados)} casos aprobados · puntaje promedio {prom}"
                 + (f" · {comp['regresiones']} regresión(es) contra la corrida anterior" if comp.get("regresiones") else ""))
    cuerpo = section("Resultado", banner(veredicto, tono) + kpi_tiles([
        ("Casos evaluados", len(evaluados), ""), ("Aprobados", aprobados, "ok" if aprobados == len(evaluados) else "warn"),
        ("Puntaje promedio", prom, "ok" if prom >= 80 else "warn" if prom >= 60 else "bad", "", "objetivo", {"pct": prom}),
        ("Regresiones", comp.get("regresiones", 0) if comp else "—", "bad" if comp.get("regresiones") else ""),
        ("Costo total (USD)", f"{costo:.2f}" if costo else "—", ""),
    ]))
    por_agente = datos.get("por_agente", {})
    filas = [[f"<code>{e(a)}</code>", v["casos"], v["aprobados"], v["puntaje"]] for a, v in sorted(por_agente.items())]
    cuerpo += section("Por agente", table(["Agente", "Casos", "Aprobados", "Puntaje promedio"], filas), ancho=6)
    color = lambda p: OK if p >= 80 else AVISO if p >= 60 else CRITICO
    cuerpo += section("Puntaje por caso", bar_chart([(c["id"], c["puntaje"], color(c["puntaje"])) for c in evaluados], 220), ancho=6)
    if comp and comp.get("casos"):
        filas = [[e(cid), v["antes"], v["despues"], e(", ".join(v["regresiones"]) or "—"), e(", ".join(v["mejoras"]) or "—")]
                 for cid, v in comp["casos"].items()]
        cuerpo += section("Comparación con la corrida anterior", table(["Caso", "Antes", "Ahora", "Regresiones", "Mejoras"], filas),
                          sub=f"Corrida anterior: {comp.get('fecha_anterior', '')}")
    for c in casos:
        m = c.get("metricas") or {}
        info = " · ".join(x for x in [
            f"Agente {c['agente']}", f"Puntaje {c['puntaje']}" if c["estado"] in ("Aprobado", "Fallido") else "",
            f"{m['turnos']} turnos" if m.get("turnos") else "", f"US$ {m['costo_usd']:.2f}" if m.get("costo_usd") else "",
            f"{m['duracion_s']} s" if m.get("duracion_s") else ""] if x)
        cuerpo_caso = f'<div class="cardbody"><p class="muted">{e(c.get("descripcion", ""))}</p><p>{estado_pill(c["estado"])} {e(info)}</p></div>'
        if c.get("motivo"):
            cuerpo_caso += doc_list([c["motivo"]])
        if c.get("chequeos"):
            cuerpo_caso += table(["Estado", "Chequeo", "Detalle"], [
                [estado_pill("Aprobado" if x["ok"] else "Fallido"), e(x["descripcion"]), e(x["detalle"])] for x in c["chequeos"]])
        if c.get("juez") and c["juez"].get("puntajes"):
            cuerpo_caso += table(["Criterio (juez, informativo)", "Puntaje", "Motivo"], [
                [e(p.get("criterio")), e(p.get("puntaje")), e(p.get("motivo"))] for p in c["juez"]["puntajes"]])
        cuerpo += section(f"Caso {c['id']}", cuerpo_caso)
    return page("Evals de los agentes", "Evals", meta([("Fecha", datos["fecha"]), ("Modelo", datos.get("modelo") or "por defecto"),
                                                       ("Casos", len(casos))]), cuerpo)


def guardar(datos, destino):
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    datos["por_agente"] = resumen_por_agente(datos["casos"])
    datos["comparacion"] = comparar(datos, corrida_anterior(destino))
    (destino / "resultados.json").write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    sys.path.insert(0, str(RAIZ / "scripts"))
    from _i18n import escribir_bilingue
    escribir_bilingue(destino / "informe-evals.html", lambda: informe_html(datos))
    return destino


# ------------------------------------------------------------------ CLI
def _sello():
    return dt.datetime.now().strftime("%Y%m%d-%H%M%S")


def main(argv=None):
    for flujo in (sys.stdout, sys.stderr):
        flujo.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Evals de los agentes de k0lmenIA")
    accion = ap.add_mutually_exclusive_group(required=True)
    accion.add_argument("--listar", action="store_true")
    accion.add_argument("--preparar", action="store_true")
    accion.add_argument("--correr", action="store_true")
    accion.add_argument("--verificar", action="store_true")
    accion.add_argument("--solo-verificar", metavar="CARPETA")
    ap.add_argument("--caso")
    ap.add_argument("--agente")
    ap.add_argument("--trabajo", help="carpeta de los workspaces (por defecto, en la carpeta temporal del sistema)")
    ap.add_argument("--resultados", default=str(RESULTADOS))
    ap.add_argument("--casos", default=str(CASOS), help=argparse.SUPPRESS)
    ap.add_argument("--modelo")
    ap.add_argument("--max-turnos", type=int, help="tope de turnos (--max-turns), si tu versión de Claude Code lo acepta")
    ap.add_argument("--presupuesto-usd", type=float, help="tope de gasto por caso (pisa el presupuesto_usd del caso; 2 por defecto)")
    ap.add_argument("--juez", action="store_true")
    a = ap.parse_args(argv)
    casos = cargar_casos(a.caso, a.agente, a.casos)
    sello = _sello()

    if a.listar:
        for c in casos:
            falta = faltantes(c)
            print(f"{c['id']:32} {c.get('agente', ''):26} {len(c.get('chequeos', []))} chequeos"
                  + (f"  (se omite: falta {', '.join(falta)})" if falta else ""))
        return 0

    if a.solo_verificar:
        datos = {"fecha": sello, "modo": "solo-verificar", "carpeta": a.solo_verificar,
                 "casos": [verificar_caso(c, a.solo_verificar) for c in casos]}
        destino = guardar(datos, Path(a.resultados) / sello)
        _imprimir(datos, destino)
        return 0 if all(c["estado"] == "Aprobado" for c in datos["casos"]) else 1

    trabajo = Path(a.trabajo) if a.trabajo else Path(tempfile.gettempdir()) / "k0lmenia-evals" / sello
    if a.verificar:
        if not a.trabajo:
            raise SystemExit("--verificar necesita --trabajo con la carpeta de la corrida")
        corrida = json.loads((trabajo / "corrida.json").read_text(encoding="utf-8")) if (trabajo / "corrida.json").exists() else {}
        datos = {"fecha": sello, "modo": "verificar", "trabajo": str(trabajo), "modelo": corrida.get("modelo"), "casos": []}
        for c in casos:
            ws = trabajo / c["id"]
            if not ws.exists():
                continue
            res = verificar_caso(c, ws)
            res["metricas"] = (corrida.get("metricas") or {}).get(c["id"])
            datos["casos"].append(res)
        destino = guardar(datos, Path(a.resultados) / sello)
        _imprimir(datos, destino)
        return 0

    corrida = {"fecha": sello, "modelo": a.modelo, "casos": [], "metricas": {}}
    datos = {"fecha": sello, "modo": "correr" if a.correr else "preparar", "trabajo": str(trabajo), "modelo": a.modelo, "casos": []}
    for c in casos:
        falta = faltantes(c)
        if falta:
            datos["casos"].append({"id": c["id"], "agente": c.get("agente", ""), "descripcion": c.get("descripcion", ""),
                                   "puntaje": 0, "estado": "N/A", "chequeos": [], "motivo": "Se omite: falta " + ", ".join(falta)})
            continue
        ws = preparar(c, trabajo / c["id"])
        corrida["casos"].append(c["id"])
        if a.preparar:
            print(f"{c['id']}: {ws}")
            print("   cd \"" + str(ws) + "\" && " + " ".join(_citar(x) for x in comando_claude(c, a.modelo, a.max_turnos, presupuesto_usd=a.presupuesto_usd)))
            continue
        print(f"Corriendo {c['id']} ({c.get('agente')})…", flush=True)
        metricas = correr_claude(c, ws, a.modelo, a.max_turnos, presupuesto_usd=a.presupuesto_usd)
        corrida["metricas"][c["id"]] = metricas
        if metricas.get("error") and not any((ws / "output").rglob("*.*")):
            datos["casos"].append({"id": c["id"], "agente": c.get("agente", ""), "descripcion": c.get("descripcion", ""),
                                   "puntaje": 0, "estado": "Bloqueado", "chequeos": [], "metricas": metricas,
                                   "motivo": f"No se pudo correr: {metricas['error']}"})
            continue
        res = verificar_caso(c, ws)
        res["metricas"] = metricas
        if a.juez:
            res["juez"] = juez(c, ws, a.modelo)
        datos["casos"].append(res)
    trabajo.mkdir(parents=True, exist_ok=True)
    (trabajo / "corrida.json").write_text(json.dumps(corrida, ensure_ascii=False, indent=2), encoding="utf-8")
    if a.preparar:
        print(f"\nWorkspaces en {trabajo}. Después de correr los pedidos: python scripts/evaluar_agentes.py --verificar --trabajo \"{trabajo}\"")
        return 0
    destino = guardar(datos, Path(a.resultados) / sello)
    _imprimir(datos, destino)
    return 0


def _citar(x):
    return x if re.fullmatch(r"[\w@%+=:,./*()-]+", x) else '"' + x.replace('"', '\\"') + '"'


def _imprimir(datos, destino):
    for c in datos["casos"]:
        print(f"  {c['estado']:10} {c['puntaje']:>3}  {c['id']}  ({c['agente']})")
        for x in c.get("chequeos", []):
            if not x["ok"]:
                print(f"             ✕ {x['descripcion']}: {x['detalle']}")
    comp = datos.get("comparacion")
    if comp:
        print(f"Contra la corrida anterior: {comp['regresiones']} regresión(es), {comp['mejoras']} mejora(s).")
    print(f"Informe: {Path(destino) / 'informe-evals.html'}")


if __name__ == "__main__":
    sys.exit(main())
