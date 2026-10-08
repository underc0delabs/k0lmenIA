#!/usr/bin/env python3
"""Genera los casos de prueba a partir de un JSON: la planilla Excel (.xlsx) y,
en el mismo paso, una versión en Markdown (.md) con los casos en tabla ASCII alineada.

Uso:
    python scripts/generar_casos.py <entrada.json> <salida.xlsx> [--limpiar]

Escribe <salida.xlsx> y, al lado, <salida.md> (misma ruta, extensión .md).
Con --limpiar, borra el JSON de entrada al terminar (útil para los temporales).

Formato del JSON de entrada:
{
  "proyecto": "NOMBRE DEL PROYECTO",
  "modulo": "HU-001 — Inicio de sesión",
  "casos": [
    {"id": "CP-001", "titulo": "...", "descripcion": "...", "precondiciones": "...",
     "datos": "Email: ...\nContraseña: ...", "pasos": "1. ...\n2. ...", "resultado": "...",
     "estado": "Pendiente", "prioridad": "Alta", "etiquetas": "@login",
     "evidencia": "", "fecha_ejecucion": "", "comentarios": "",
     "falta_info": ["FI-01"]}
  ],
  "falta_informacion": [
    {"id": "FI-01", "que_falta": "Texto del mensaje al bloquearse la cuenta",
     "donde_se_busco": "Historia, comentarios, Figma", "impacto": "CA3",
     "pregunta": "¿Cuál es el mensaje exacto?", "estado": "Abierto"}
  ]
}

"falta_info" (opcional, por caso): IDs de "Falta información" de los que depende el caso.
Esos casos se resaltan en la planilla, suman la etiqueta @falta-info y el detalle en
Comentarios. "falta_informacion" (opcional) arma la hoja "Falta información" del .xlsx
y la sección del mismo nombre en el .md (ver skill investigacion-contexto).
"""
import sys
import os
import re
import json
import textwrap
import subprocess

COLS = ["ID #", "Título", "Descripción", "Precondiciones", "Datos de Prueba", "Pasos",
        "Resultado Esperado", "Estado", "Prioridad", "#Etiquetas", "Evidencia",
        "Fecha de Ejecución", "Comentarios"]
KEYS = ["id", "titulo", "descripcion", "precondiciones", "datos", "pasos", "resultado",
        "estado", "prioridad", "etiquetas", "evidencia", "fecha_ejecucion", "comentarios"]
# El .md tiene: (1) una tabla Resumen y (2) la tabla con todos los casos (una fila por caso).
# Resumen: 4 columnas, ordenado por prioridad (lo más crítico primero).
RESUMEN_COLS = ["ID", "Título", "Resultado esperado", "Prioridad"]
RESUMEN_KEYS = ["id", "titulo", "resultado", "prioridad"]
RESUMEN_W = [6, 24, 32, 10]

# Tabla de casos del .md. Quita Estado (siempre "Pendiente"), Evidencia, Fecha de Ejecución
# y Comentarios para no ensanchar de más; todas esas columnas sí están en el .xlsx.
MD_COLS = ["ID #", "Título", "Descripción", "Precondiciones", "Datos de Prueba", "Pasos",
           "Resultado Esperado", "Prioridad", "#Etiquetas"]
MD_KEYS = ["id", "titulo", "descripcion", "precondiciones", "datos", "pasos",
           "resultado", "prioridad", "etiquetas"]
# Ancho máximo por columna (las celdas largas se ajustan a varias líneas dentro de la celda).
MD_MAXW = [6, 16, 22, 16, 18, 24, 20, 9, 12]

FI_COLS = ["ID", "Qué falta", "Dónde se buscó", "Impacto", "Pregunta para el PO", "Estado"]
FI_KEYS = ["id", "que_falta", "donde_se_busco", "impacto", "pregunta", "estado"]
FI_W = [8, 40, 30, 22, 40, 11]


def _ids_fi(caso):
    v = caso.get("falta_info") or []
    if isinstance(v, str):
        v = [x for x in v.replace(";", ",").split(",")]
    return [str(x).strip() for x in v if str(x).strip()]


def preparar_falta_info(data):
    """Marca los casos que dependen de información faltante: etiqueta @falta-info y
    el detalle de cada FI en Comentarios (sin duplicar si ya estaba)."""
    detalle = {str(f.get("id", "")).strip(): f for f in data.get("falta_informacion", []) or []}
    for caso in data.get("casos", []):
        ids = _ids_fi(caso)
        if not ids:
            continue
        etiquetas = str(caso.get("etiquetas", "") or "")
        if "@falta-info" not in etiquetas:
            caso["etiquetas"] = (etiquetas + " @falta-info").strip()
        comentarios = str(caso.get("comentarios", "") or "")
        notas = []
        for fid in ids:
            if fid in comentarios:
                continue
            que = (detalle.get(fid) or {}).get("que_falta", "")
            notas.append(f"FALTA INFORMACIÓN ({fid}){': ' + que if que else ''}")
        if notas:
            caso["comentarios"] = "\n".join(notas + ([comentarios] if comentarios else []))
        caso["_fi"] = ", ".join(ids)


# Orden de prioridad: más crítico primero, prioridad baja al final.
PRIORIDAD_ORDEN = {"crítica": 0, "critica": 0, "alta": 1, "media": 2, "baja": 3}


def _need(pkg, instalar=None):
    """Asegura que la librería esté disponible. Si falta, intenta instalarla; si no
    puede, avisa con un mensaje claro y corta (en vez de fallar con un traceback)."""
    instalar = instalar or pkg
    try:
        __import__(pkg)
        return
    except ImportError:
        pass
    subprocess.run([sys.executable, "-m", "pip", "install", instalar, "--quiet"], check=False)
    try:
        __import__(pkg)
    except ImportError:
        sys.exit(
            f"\n[!] Falta la librería '{instalar}' y no se pudo instalar sola.\n"
            f"    Instalá las dependencias con:  pip install -r requirements.txt\n"
            f"    (o directamente:  pip install {instalar})\n")


def to_xlsx(data, out):
    _need("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    proyecto = data.get("proyecto", "NOMBRE DEL PROYECTO")
    modulo = data.get("modulo", "Módulo")
    casos = data.get("casos", [])
    widths = [6, 30, 42, 32, 30, 46, 40, 13, 11, 22, 20, 18, 30]
    n = len(COLS)
    last_col = chr(ord("A") + n - 1)

    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[:\\/?*\[\]]", "-", modulo)[:31]) or "Casos"
    gris = PatternFill("solid", fgColor="404040")
    azul = PatternFill("solid", fgColor="4472C4")
    blanco_bold = Font(name="Arial", color="FFFFFF", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = proyecto
    ws["A1"].fill = gris
    ws["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"Módulo: {modulo}"
    ws["A2"].fill = azul
    ws["A2"].font = blanco_bold
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    for i, c in enumerate(COLS):
        cell = ws.cell(row=3, column=i + 1, value=c)
        cell.fill = azul
        cell.font = blanco_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = borde
    ws.row_dimensions[3].height = 24
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w

    data_font = Font(name="Arial", size=10)
    ambar = PatternFill("solid", fgColor="FFF2CC")
    ambar_fuerte = Font(name="Arial", size=10, bold=True, color="9C5700")
    start = 4
    for r, caso in enumerate(casos):
        row = start + r
        falta = bool(caso.get("_fi"))
        for c, key in enumerate(KEYS):
            cell = ws.cell(row=row, column=c + 1, value=caso.get(key, ""))
            cell.font = ambar_fuerte if (falta and key == "comentarios") else data_font
            cell.border = borde
            if falta:
                cell.fill = ambar
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if c in (0, 7, 8) else "left")
        ws.row_dimensions[row].height = 95

    last_row = start + len(casos) + 50
    dv_estado = DataValidation(type="list",
        formula1='"N/A,Pendiente,En ejecución,Aprobado,Fallido,Bloqueado"', allow_blank=True)
    dv_prio = DataValidation(type="list", formula1='"Crítica,Alta,Media,Baja"', allow_blank=True)
    ws.add_data_validation(dv_estado)
    ws.add_data_validation(dv_prio)
    dv_estado.add(f"H{start}:H{last_row}")
    dv_prio.add(f"I{start}:I{last_row}")
    ws.freeze_panes = "A4"

    # Hoja "Falta información": lo que no se encontró en ninguna fuente (skill investigacion-contexto).
    faltantes = data.get("falta_informacion", []) or []
    if faltantes:
        wf = wb.create_sheet("Falta información")
        wf.merge_cells("A1:F1")
        wf["A1"] = f"Falta información — {modulo}"
        wf["A1"].fill = PatternFill("solid", fgColor="9C5700")
        wf["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=13)
        wf["A1"].alignment = Alignment(horizontal="center", vertical="center")
        wf.row_dimensions[1].height = 24
        for i, c in enumerate(FI_COLS):
            cell = wf.cell(row=2, column=i + 1, value=c)
            cell.fill = azul
            cell.font = blanco_bold
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = borde
            wf.column_dimensions[chr(ord("A") + i)].width = FI_W[i] + 6
        for r, f in enumerate(faltantes):
            for c, key in enumerate(FI_KEYS):
                cell = wf.cell(row=3 + r, column=c + 1, value=f.get(key, "Abierto" if key == "estado" else ""))
                cell.font = data_font
                cell.border = borde
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            wf.row_dimensions[3 + r].height = 60
        afectados = sum(1 for c in casos if c.get("_fi"))
        nota = ws.cell(row=start + len(casos) + 1, column=1,
                       value=f"Casos resaltados: dependen de información faltante ({afectados}). Ver la hoja \"Falta información\".")
        nota.font = Font(name="Arial", size=9, italic=True, color="9C5700")
    wb.save(out)


def _wrap(texto, ancho):
    """Ajusta el texto a 'ancho' caracteres, respetando los saltos de línea existentes."""
    lineas = []
    for linea in str(texto).split("\n"):
        if linea.strip() == "":
            lineas.append("")
        else:
            lineas.extend(textwrap.wrap(linea, width=ancho) or [""])
    return "\n".join(lineas)


def to_md(data, out):
    _need("tabulate")
    from tabulate import tabulate
    casos = data.get("casos", [])
    titulo = (f"Casos de prueba — {data.get('modulo', '')}").strip(" —")
    partes = [f"# {titulo}", "", "## Resumen", ""]
    filas_r = [[_wrap(c.get(k, ""), w) for k, w in zip(RESUMEN_KEYS, RESUMEN_W)] for c in casos]
    cols_r = RESUMEN_COLS
    if any(c.get("_fi") for c in casos):
        cols_r = RESUMEN_COLS + ["Falta info"]
        filas_r = [f + [_wrap(c.get("_fi", ""), 10)] for f, c in zip(filas_r, casos)]
    partes += ["```", tabulate(filas_r, headers=cols_r, tablefmt="grid"), "```", ""]
    partes += ["## Detalle de los casos", ""]
    filas = [[c.get(k, "") for k in MD_KEYS] for c in casos]
    tabla = tabulate(filas, headers=MD_COLS, tablefmt="grid", maxcolwidths=MD_MAXW)
    partes += ["```", tabla, "```", ""]
    faltantes = data.get("falta_informacion", []) or []
    if faltantes:
        afectados = [f"{c.get('id', '')} ({c['_fi']})" for c in casos if c.get("_fi")]
        partes += ["## Falta información", "",
                   "Datos que no aparecen en ninguna fuente consultada. Los casos que dependen de ellos "
                   "llevan la etiqueta `@falta-info`" + (f": {', '.join(afectados)}." if afectados else "."), ""]
        filas_fi = [[f.get(k, "Abierto" if k == "estado" else "") for k in FI_KEYS] for f in faltantes]
        partes += ["```", tabulate(filas_fi, headers=FI_COLS, tablefmt="grid", maxcolwidths=FI_W), "```", ""]
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(partes) + "\n")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = [a for a in sys.argv[1:] if a.startswith("--")]
    if len(args) < 2:
        print("Uso: python scripts/generar_casos.py <entrada.json> <salida.xlsx> [--limpiar]")
        sys.exit(1)
    entrada, salida = args[0], args[1]
    if not salida.lower().endswith(".xlsx"):
        print("La salida debe terminar en .xlsx")
        sys.exit(1)
    with open(entrada, encoding="utf-8") as f:
        data = json.load(f)
    # Ordenar los casos por prioridad: los más críticos primero, los de prioridad baja al final.
    # El orden es estable: dentro de la misma prioridad, se respeta el orden original.
    data["casos"] = sorted(
        data.get("casos", []),
        key=lambda c: PRIORIDAD_ORDEN.get(str(c.get("prioridad", "")).strip().lower(), 99))
    preparar_falta_info(data)
    to_xlsx(data, salida)
    print(f"OK: {salida}")
    md = re.sub(r"\.xlsx$", ".md", salida, flags=re.IGNORECASE)
    to_md(data, md)
    print(f"OK: {md}")
    # Con --limpiar, borra el JSON de entrada (cross-platform, sin depender de 'rm').
    if "--limpiar" in flags:
        try:
            os.remove(entrada)
            print(f"Borrado temporal: {entrada}")
        except OSError:
            pass


if __name__ == "__main__":
    main()
