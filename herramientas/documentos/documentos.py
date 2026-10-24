#!/usr/bin/env python3
"""
Relevamiento de documentos generados (PDF y Excel/CSV), de solo lectura, para validarlos en QA.

Uso (desde la raíz del repo):
    python herramientas/documentos/documentos.py --archivo output/.../factura.pdf [--salida ...]
    python herramientas/documentos/documentos.py --archivo reporte.xlsx

Qué hace:
- PDF (pypdf): extrae texto por página, metadata, cantidad de páginas y si está cifrado.
- Excel (openpyxl): lista hojas, dimensiones y una muestra de celdas; CSV: filas de muestra.
- Chequea **placeholders sin reemplazar** (variables de plantilla que quedaron, Lorem ipsum, TODO) — el
  bug más común de un documento generado.

No modifica el documento. La verificación de que los datos/cálculos son correctos y el layout está bien la
hace el agente probador-documentos con este insumo y los datos esperados.
"""
import argparse
import json
import re
import sys
from pathlib import Path

PLACEHOLDERS = [
    ("Variable de plantilla", re.compile(r"\{\{\s*[\w.]+\s*\}\}|\$\{[\w.]+\}|%\(?[\w]+\)?[sd]|<[\w]+>")),
    ("Lorem ipsum", re.compile(r"(?i)lorem ipsum")),
    ("Marcador pendiente", re.compile(r"\b(TODO|FIXME|XXX|TBD|PENDIENTE|PLACEHOLDER|CAMBIAR)\b")),
    ("Dato de ejemplo", re.compile(r"(?i)\b(xxx+|aaa+|test test|dummy|sample text|texto de ejemplo)\b")),
]


def buscar_placeholders(texto):
    """[(tipo, fragmento)] de marcadores que no deberían quedar en un documento final. Puro y testeable."""
    out = []
    for tipo, rx in PLACEHOLDERS:
        for m in rx.finditer(texto or ""):
            out.append((tipo, m.group(0)[:40]))
    return out


def leer_pdf(ruta):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("Falta pypdf: instalalo con `pip install pypdf`.")
    r = PdfReader(ruta)
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            return {"tipo": "pdf", "cifrado": True, "paginas": len(r.pages), "texto": "", "metadata": {}}
    texto = []
    for pg in r.pages:
        try:
            texto.append(pg.extract_text() or "")
        except Exception:
            texto.append("")
    meta = {}
    try:
        meta = {k.lstrip("/"): str(v) for k, v in (r.metadata or {}).items()}
    except Exception:
        pass
    full = "\n".join(texto)
    return {"tipo": "pdf", "cifrado": bool(r.is_encrypted), "paginas": len(r.pages), "texto": full,
            "metadata": meta, "placeholders": buscar_placeholders(full)}


def leer_xlsx(ruta, max_filas=50):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Falta openpyxl: instalalo con `pip install openpyxl`.")
    wb = load_workbook(ruta, read_only=True, data_only=True)
    hojas, texto = {}, []
    for hoja in wb.worksheets:
        filas = []
        for i, fila in enumerate(hoja.iter_rows(values_only=True)):
            if i >= max_filas:
                break
            filas.append([("" if c is None else str(c)) for c in fila])
            texto += [c for c in filas[-1] if c]
        hojas[hoja.title] = {"dimension": hoja.dimensions, "muestra": filas}
    full = " ".join(texto)
    return {"tipo": "xlsx", "hojas": hojas, "placeholders": buscar_placeholders(full)}


def leer_csv(ruta, max_filas=50):
    import csv
    filas, texto = [], []
    with open(ruta, encoding="utf-8", errors="replace", newline="") as f:
        for i, fila in enumerate(csv.reader(f)):
            if i >= max_filas:
                break
            filas.append(fila)
            texto += fila
    full = " ".join(texto)
    return {"tipo": "csv", "filas": filas, "placeholders": buscar_placeholders(full)}


def relevar(ruta):
    ext = Path(ruta).suffix.lower()
    if ext == ".pdf":
        return leer_pdf(ruta)
    if ext in (".xlsx", ".xlsm"):
        return leer_xlsx(ruta)
    if ext in (".csv", ".tsv"):
        return leer_csv(ruta)
    raise RuntimeError(f"Extensión no soportada: {ext} (pdf, xlsx, csv).")


def main(argv=None):
    for flujo in (sys.stdout, sys.stderr):
        try:
            flujo.reconfigure(encoding="utf-8")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="Relevamiento de documentos generados (PDF/Excel/CSV), solo lectura.")
    ap.add_argument("--archivo", required=True)
    ap.add_argument("--salida", default="")
    args = ap.parse_args(argv)
    if not Path(args.archivo).exists():
        print(f"[documentos] No existe: {args.archivo}", file=sys.stderr)
        return 2
    try:
        inv = relevar(args.archivo)
    except RuntimeError as ex:
        print(f"[documentos] {ex}", file=sys.stderr)
        return 2
    inv["archivo"] = args.archivo
    texto = json.dumps(inv, ensure_ascii=False, indent=2)
    if args.salida:
        p = Path(args.salida)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(texto, encoding="utf-8")
        print(f"[documentos] Datos en {p}", file=sys.stderr)
    print(texto)
    ph = inv.get("placeholders") or []
    print(f"[documentos] {inv['tipo']} · {len(ph)} placeholder(s) sin reemplazar.", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
